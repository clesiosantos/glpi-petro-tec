#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Importação da Base 4711 para GLPI 10.

Origem:
  Planilha: Base_dedados_4711_com_Senior.xlsx (na raiz do repositório)
  Aba:      Posto de Trabalho e Fiscais

Regras implementadas:
- Coluna B (GERENCIA LOTACAO):
    * cria/garante um Grupo GLPI na entidade G4F
    * associa o usuário ao grupo
    * cria/garante o valor no campo dinâmico "Gerencia Lotação"
- Coluna C (LOCAL):
    * cria/garante o valor no campo dinâmico "Localização fisica Gerencia"
- Coluna D (STATUS DA MOBILIZACAO):
    * cria/garante o valor no campo dinâmico "Status Mobilização"
- Coluna J (PERFIL PADRÃO):
    * associa o perfil indicado ao usuário na entidade G4F, recursivo
    * define esse perfil como padrão do usuário
- Usuário:
    * login = Matrícula Senior
    * para usuários já implantados, o login é migrado para a Matrícula Senior preservando o users_id
    * primeiro nome / último nome são derivados de NOME DO TECNICO
    * comentário = "ITEM PPU: <valor> :: PREPOSTO: <valor>"
    * grupo padrão = grupo criado a partir de GERENCIA LOTACAO
    * senha inicial padrão = variável GLPI_DEFAULT_PASSWORD
- Linhas cujo NOME DO TECNICO seja "não mobilizar de imediato" são IGNORADAS.

Campos dinâmicos que NÃO serão preenchidos nesta primeira carga:
- Preposto
- Chave Colaborador
- Líder

Modos:
  --dry-run : valida planilha, API, perfis, container e campos; não grava nada.
  --apply   : executa a importação.

Exemplo:
  source /root/.glpi.env
  python3 scripts/import_base4711.py --dry-run
"""

import argparse
import json
import os
import re
import sys
import unicodedata
from pathlib import Path
from collections import Counter
from typing import Any, Dict, Iterable, List, Optional, Tuple

import requests

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "ERRO: módulo openpyxl não instalado.\n"
        "Instale antes com:\n"
        "  dnf -y install python3-openpyxl\n"
        "ou:\n"
        "  python3 -m pip install openpyxl",
        file=sys.stderr,
    )
    raise SystemExit(2)


DEFAULT_API_URL = "https://tec.g4f.sharksolucoes.com.br/apirest.php"
DEFAULT_XLSX = str(
    Path(__file__).resolve().parents[1] / "Base_dedados_4711_com_Senior.xlsx"
)
SHEET_NAME = "Posto de Trabalho e Fiscais"

ROOT_ENTITY_ID = 0
ADMIN_PROFILE_ID = 4  # ZZ-Super-Admin

CONTAINER_LABEL = "Agrupamento"
CONTAINER_ITEMTYPE = "User"

FIELD_LABEL_STATUS = "Status Mobilização"
FIELD_LABEL_GERENCIA = "Gerencia Lotação"
FIELD_LABEL_LOCAL = "Localização fisica Gerencia"

SKIP_NAMES = {
    "não mobilizar de imediato",
    "nao mobilizar de imediato",
}

GROUP_FLAGS = {
    "groups_id": 0,
    "entities_id": ROOT_ENTITY_ID,
    "is_recursive": 1,
    "is_requester": 1,
    "is_watcher": 0,
    "is_assign": 1,
    "is_task": 0,
    "is_notify": 0,
    "is_manager": 0,
    "is_itemgroup": 1,
    "is_usergroup": 1,
}


def load_shell_env_file(path: str = "/root/.glpi.env") -> None:
    """Carrega arquivo simples no formato export CHAVE='valor'."""
    if not os.path.exists(path):
        return

    with open(path, "r", encoding="utf-8") as fh:
        for raw in fh:
            line = raw.strip()
            if not line or line.startswith("#"):
                continue
            if line.startswith("export "):
                line = line[7:].strip()
            if "=" not in line:
                continue

            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip()

            if (
                len(value) >= 2
                and value[0] == value[-1]
                and value[0] in ("'", '"')
            ):
                value = value[1:-1]

            if key and key not in os.environ:
                os.environ[key] = value


load_shell_env_file()

API_URL = os.environ.get("GLPI_API_URL", DEFAULT_API_URL).rstrip("/")
USER_TOKEN = os.environ.get("GLPI_USER_TOKEN")
APP_TOKEN = os.environ.get("GLPI_APP_TOKEN")
DEFAULT_PASSWORD = os.environ.get("GLPI_DEFAULT_PASSWORD")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_compare(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.casefold().strip()
    value = re.sub(r"\s+", " ", value)
    return value


def excel_code(value: Any) -> str:
    """Preserva códigos como 1.1, 3.3 e 0 sem transformar 0 em 0.0."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10f}".rstrip("0").rstrip(".")
    return clean_text(value)


def split_name(full_name: str) -> Tuple[str, str]:
    parts = re.split(r"\s+", full_name.strip())
    if not parts:
        return "", ""
    firstname = parts[0]
    realname = " ".join(parts[1:]) if len(parts) > 1 else ""
    return firstname, realname


def chunked(iterable: Iterable[Any], size: int) -> Iterable[List[Any]]:
    chunk: List[Any] = []
    for item in iterable:
        chunk.append(item)
        if len(chunk) >= size:
            yield chunk
            chunk = []
    if chunk:
        yield chunk


class GLPI:
    def __init__(self):
        if not USER_TOKEN:
            raise RuntimeError("GLPI_USER_TOKEN não definido.")
        if not APP_TOKEN:
            raise RuntimeError("GLPI_APP_TOKEN não definido.")

        self.http = requests.Session()
        self.http.headers.update(
            {
                "App-Token": APP_TOKEN,
                "Content-Type": "application/json",
            }
        )
        self.session_token: Optional[str] = None

    def init(self) -> None:
        r = self.http.get(
            f"{API_URL}/initSession",
            headers={
                "Authorization": f"user_token {USER_TOKEN}",
                "App-Token": APP_TOKEN,
                "Content-Type": "application/json",
            },
            timeout=30,
        )
        if r.status_code != 200:
            raise RuntimeError(
                f"initSession: HTTP {r.status_code}: {r.text}"
            )

        token = r.json().get("session_token")
        if not token:
            raise RuntimeError(f"session_token não retornado: {r.text}")

        self.session_token = token
        self.http.headers["Session-Token"] = token

        r = self.http.post(
            f"{API_URL}/changeActiveProfile",
            json={"profiles_id": ADMIN_PROFILE_ID},
            timeout=30,
        )
        if r.status_code != 200 or r.text.strip().lower() != "true":
            raise RuntimeError(
                f"changeActiveProfile: HTTP {r.status_code}: {r.text}"
            )

        r = self.http.post(
            f"{API_URL}/changeActiveEntities",
            json={
                "entities_id": ROOT_ENTITY_ID,
                "is_recursive": True,
            },
            timeout=30,
        )
        if r.status_code != 200 or r.text.strip().lower() != "true":
            raise RuntimeError(
                f"changeActiveEntities: HTTP {r.status_code}: {r.text}"
            )

        print("OK - sessão GLPI criada")
        print("OK - perfil ZZ-Super-Admin ativo")
        print("OK - entidade G4F ativa recursivamente")

    def close(self) -> None:
        if not self.session_token:
            return
        try:
            self.http.get(f"{API_URL}/killSession", timeout=15)
        except Exception:
            pass

    def get_all(
        self,
        itemtype: str,
        *,
        page_size: int = 200,
        allow_missing: bool = False,
    ) -> List[Dict[str, Any]]:
        result: List[Dict[str, Any]] = []
        start = 0

        while True:
            end = start + page_size - 1
            r = self.http.get(
                f"{API_URL}/{itemtype}/",
                params={
                    "range": f"{start}-{end}",
                    "get_hateoas": "false",
                },
                timeout=60,
            )

            if allow_missing and r.status_code in (400, 404):
                return []

            if r.status_code not in (200, 206):
                raise RuntimeError(
                    f"GET {itemtype}: HTTP {r.status_code}: {r.text}"
                )

            data = r.json()
            if not isinstance(data, list):
                raise RuntimeError(
                    f"Resposta inesperada em {itemtype}: {data}"
                )

            result.extend(data)

            if len(data) < page_size:
                break
            start += page_size

        return result

    def add(self, itemtype: str, payload: Dict[str, Any]) -> int:
        r = self.http.post(
            f"{API_URL}/{itemtype}/",
            json={"input": payload},
            timeout=60,
        )

        if r.status_code not in (200, 201):
            raise RuntimeError(
                f"POST {itemtype}: HTTP {r.status_code}: {r.text}\n"
                f"Payload: {json.dumps(payload, ensure_ascii=False)}"
            )

        data = r.json()
        if isinstance(data, dict) and data.get("id") is not None:
            return int(data["id"])

        if isinstance(data, list) and data and data[0].get("id") is not None:
            return int(data[0]["id"])

        raise RuntimeError(
            f"POST {itemtype}: ID não retornado. Resposta: {data}"
        )

    def update(
        self,
        itemtype: str,
        item_id: int,
        payload: Dict[str, Any],
    ) -> None:
        body = dict(payload)
        body["id"] = int(item_id)

        r = self.http.put(
            f"{API_URL}/{itemtype}/{item_id}",
            json={"input": body},
            timeout=60,
        )

        if r.status_code not in (200, 201):
            raise RuntimeError(
                f"PUT {itemtype}/{item_id}: "
                f"HTTP {r.status_code}: {r.text}\n"
                f"Payload: {json.dumps(body, ensure_ascii=False)}"
            )



def normalized_header_index(headers: List[Any]) -> Dict[str, int]:
    return {
        normalize_compare(clean_text(value)): idx
        for idx, value in enumerate(headers)
        if clean_text(value)
    }


def read_rows(
    xlsx_path: str,
) -> Tuple[List[Dict[str, str]], List[int]]:
    if not os.path.exists(xlsx_path):
        raise RuntimeError(f"Planilha não encontrada: {xlsx_path}")

    book = load_workbook(xlsx_path, read_only=True, data_only=True)

    if SHEET_NAME not in book.sheetnames:
        book.close()
        raise RuntimeError(
            f"Aba '{SHEET_NAME}' não encontrada. "
            f"Abas disponíveis: {book.sheetnames}"
        )

    ws = book[SHEET_NAME]

    headers = [clean_text(cell.value) for cell in ws[1]]
    index = {h: i for i, h in enumerate(headers)}
    normalized_index = normalized_header_index(headers)

    required = [
        "NOME DO TECNICO",
        "GERENCIA LOTACAO",
        "LOCAL",
        "STATUS DA MOBILIZACAO",
        "PREPOSTO",
        "ITEM PPU",
        "PERFIL PADRÃO",
    ]

    missing = [name for name in required if name not in index]
    if missing:
        book.close()
        raise RuntimeError(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}"
        )

    senior_col_idx = normalized_index.get(
        normalize_compare("Matricula Senior")
    )

    if senior_col_idx is None:
        book.close()
        raise RuntimeError(
            "Coluna obrigatória ausente: Matricula Senior"
        )

    rows: List[Dict[str, str]] = []
    skipped: List[int] = []

    for excel_row, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        name = clean_text(values[index["NOME DO TECNICO"]])

        if not name:
            continue

        if normalize_compare(name) in {
            normalize_compare(x) for x in SKIP_NAMES
        }:
            skipped.append(excel_row)
            continue

        senior_login = excel_code(values[senior_col_idx])

        if not senior_login:
            book.close()
            raise RuntimeError(
                f"Linha {excel_row}: Matrícula Senior vazia "
                f"para '{name}'."
            )

        row = {
            "excel_row": str(excel_row),
            "full_name": name,
            "login": senior_login,
            "gerencia": clean_text(
                values[index["GERENCIA LOTACAO"]]
            ),
            "local": clean_text(
                values[index["LOCAL"]]
            ),
            "status": clean_text(
                values[index["STATUS DA MOBILIZACAO"]]
            ),
            "preposto": clean_text(
                values[index["PREPOSTO"]]
            ),
            "ppu": excel_code(
                values[index["ITEM PPU"]]
            ),
            "profile": clean_text(
                values[index["PERFIL PADRÃO"]]
            ),
        }

        for key in ("login", "gerencia", "local", "status", "profile"):
            if not row[key]:
                book.close()
                raise RuntimeError(
                    f"Linha {excel_row}: campo obrigatório vazio: {key}"
                )

        rows.append(row)

    book.close()

    login_counter = Counter(normalize_compare(r["login"]) for r in rows)
    duplicated_logins = sorted(
        login for login, count in login_counter.items() if count > 1
    )
    if duplicated_logins:
        raise RuntimeError(
            "Matrículas Senior duplicadas na carga: "
            + ", ".join(duplicated_logins)
        )

    return rows, skipped

def find_container_and_fields(
    glpi: GLPI,
) -> Tuple[
    Dict[str, Any],
    Dict[str, Dict[str, Any]],
]:
    containers = glpi.get_all("PluginFieldsContainer")
    fields = glpi.get_all("PluginFieldsField")

    target_container: Optional[Dict[str, Any]] = None

    for c in containers:
        label = clean_text(c.get("label"))
        itemtypes = c.get("itemtypes")

        try:
            if isinstance(itemtypes, str):
                parsed = json.loads(itemtypes)
            elif isinstance(itemtypes, list):
                parsed = itemtypes
            else:
                parsed = []
        except Exception:
            parsed = []

        if (
            normalize_compare(label)
            == normalize_compare(CONTAINER_LABEL)
            and CONTAINER_ITEMTYPE in parsed
        ):
            target_container = c
            break

    if not target_container:
        raise RuntimeError(
            f"Container '{CONTAINER_LABEL}' para User não localizado."
        )

    container_id = int(target_container["id"])

    selected: Dict[str, Dict[str, Any]] = {}

    wanted = {
        normalize_compare(FIELD_LABEL_STATUS): "status",
        normalize_compare(FIELD_LABEL_GERENCIA): "gerencia",
        normalize_compare(FIELD_LABEL_LOCAL): "local",
    }

    for field in fields:
        if int(field.get("plugin_fields_containers_id") or 0) != container_id:
            continue

        label_norm = normalize_compare(clean_text(field.get("label")))
        key = wanted.get(label_norm)

        if key:
            selected[key] = field

    missing = [
        label
        for key, label in [
            ("status", FIELD_LABEL_STATUS),
            ("gerencia", FIELD_LABEL_GERENCIA),
            ("local", FIELD_LABEL_LOCAL),
        ]
        if key not in selected
    ]

    if missing:
        raise RuntimeError(
            "Campos dinâmicos não localizados no container "
            f"'{CONTAINER_LABEL}': {', '.join(missing)}"
        )

    for key, field in selected.items():
        if clean_text(field.get("type")) != "dropdown":
            raise RuntimeError(
                f"Campo '{field.get('label')}' não é do tipo dropdown. "
                f"Tipo encontrado: {field.get('type')}"
            )

    return target_container, selected


def dropdown_itemtype(field_internal_name: str) -> str:
    if not field_internal_name:
        raise RuntimeError("Nome interno de campo vazio.")
    return (
        "PluginFields"
        + field_internal_name[0].upper()
        + field_internal_name[1:]
        + "Dropdown"
    )


def container_instance_itemtype(container_name: str) -> str:
    # PluginFieldsContainer::getClassname('User', <container_name>)
    name = re.sub(r"s$", "", container_name, flags=re.IGNORECASE)
    system_name = ("User" + name).lower()
    return "PluginFields" + system_name[0].upper() + system_name[1:]


def build_existing_user_maps(
    users: List[Dict[str, Any]],
) -> Tuple[
    Dict[str, Dict[str, Any]],
    Dict[str, List[Dict[str, Any]]],
]:
    by_login: Dict[str, Dict[str, Any]] = {}
    by_fullname: Dict[str, List[Dict[str, Any]]] = {}

    for u in users:
        login = normalize_compare(clean_text(u.get("name")))
        if login:
            by_login[login] = u

        firstname = clean_text(u.get("firstname"))
        realname = clean_text(u.get("realname"))
        fullname = normalize_compare(
            " ".join(x for x in [firstname, realname] if x)
        )

        if fullname:
            by_fullname.setdefault(fullname, []).append(u)

    return by_login, by_fullname

def resolve_user_for_row(
    row: Dict[str, str],
    users_by_login: Dict[str, Dict[str, Any]],
    users_by_fullname: Dict[str, List[Dict[str, Any]]],
) -> Optional[Dict[str, Any]]:
    """Localiza por nome ou Matrícula Senior e bloqueia colisões."""
    fullname_key = normalize_compare(row["full_name"])
    desired_login_key = normalize_compare(row["login"])

    matches = users_by_fullname.get(fullname_key, [])
    if len(matches) > 1:
        raise RuntimeError(
            f"Mais de um usuário existente com o nome '{row['full_name']}'."
        )

    by_name = matches[0] if matches else None
    by_login = users_by_login.get(desired_login_key)

    if by_name and by_login:
        if int(by_name["id"]) != int(by_login["id"]):
            raise RuntimeError(
                f"Conflito de login: Matrícula Senior {row['login']} já pertence "
                f"ao usuário ID={by_login.get('id')} login='{by_login.get('name')}', "
                f"mas o nome '{row['full_name']}' corresponde ao "
                f"usuário ID={by_name.get('id')}."
            )
        return by_name

    if by_name:
        return by_name

    if by_login:
        existing_fullname = normalize_compare(
            " ".join(
                x for x in [
                    clean_text(by_login.get("firstname")),
                    clean_text(by_login.get("realname")),
                ]
                if x
            )
        )
        if existing_fullname and existing_fullname != fullname_key:
            raise RuntimeError(
                f"Conflito de login: Matrícula Senior {row['login']} já está "
                f"em uso pelo usuário ID={by_login.get('id')} "
                f"('{existing_fullname}')."
            )
        return by_login

    return None



def get_profiles_by_name(
    profiles: List[Dict[str, Any]],
) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for p in profiles:
        name = clean_text(p.get("name"))
        if name:
            result[normalize_compare(name)] = int(p["id"])
    return result


def map_existing_groups(
    groups: List[Dict[str, Any]],
) -> Dict[str, Dict[str, Any]]:
    result: Dict[str, Dict[str, Any]] = {}
    for g in groups:
        if int(g.get("entities_id") or 0) != ROOT_ENTITY_ID:
            continue
        name = clean_text(g.get("name"))
        if name:
            result[normalize_compare(name)] = g
    return result


def map_dropdown(
    rows: List[Dict[str, Any]],
) -> Dict[str, Dict[str, Any]]:
    result: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        name = clean_text(row.get("name"))
        if name:
            result[normalize_compare(name)] = row
    return result


def ensure_dropdown_value(
    glpi: GLPI,
    itemtype: str,
    existing_map: Dict[str, Dict[str, Any]],
    name: str,
) -> int:
    key = normalize_compare(name)

    if key in existing_map:
        return int(existing_map[key]["id"])

    item_id = glpi.add(
        itemtype,
        {
            "name": name,
            "entities_id": ROOT_ENTITY_ID,
            "is_recursive": 1,
        },
    )

    existing_map[key] = {
        "id": item_id,
        "name": name,
        "entities_id": ROOT_ENTITY_ID,
        "is_recursive": 1,
    }

    print(f"CRIADO dropdown {itemtype}: ID={item_id} {name}")
    return item_id


def ensure_group(
    glpi: GLPI,
    groups_map: Dict[str, Dict[str, Any]],
    group_name: str,
) -> int:
    key = normalize_compare(group_name)

    if key in groups_map:
        return int(groups_map[key]["id"])

    payload = dict(GROUP_FLAGS)
    payload["name"] = group_name
    payload["comment"] = "Importação Base 4711"

    group_id = glpi.add("Group", payload)

    groups_map[key] = {
        "id": group_id,
        **payload,
    }

    print(f"CRIADO grupo: ID={group_id} {group_name}")
    return group_id


def ensure_profile_user(
    glpi: GLPI,
    existing: set,
    user_id: int,
    profile_id: int,
) -> None:
    key = (user_id, profile_id, ROOT_ENTITY_ID, 1)

    if key in existing:
        return

    glpi.add(
        "Profile_User",
        {
            "users_id": user_id,
            "profiles_id": profile_id,
            "entities_id": ROOT_ENTITY_ID,
            "is_recursive": 1,
            "is_dynamic": 0,
        },
    )

    existing.add(key)


def ensure_group_user(
    glpi: GLPI,
    existing: set,
    user_id: int,
    group_id: int,
) -> None:
    key = (user_id, group_id)

    if key in existing:
        return

    glpi.add(
        "Group_User",
        {
            "users_id": user_id,
            "groups_id": group_id,
            "is_dynamic": 0,
        },
    )

    existing.add(key)


def dry_run_report(
    rows: List[Dict[str, str]],
    skipped: List[int],
    profiles_by_name: Dict[str, int],
    groups_map: Dict[str, Dict[str, Any]],
    users: List[Dict[str, Any]],
    dropdown_maps: Dict[str, Dict[str, Dict[str, Any]]],
    container: Dict[str, Any],
    fields: Dict[str, Dict[str, Any]],
    instance_itemtype: str,
    instance_api_rows: List[Dict[str, Any]],
) -> None:
    gerencias = sorted({r["gerencia"] for r in rows})
    locais = sorted({r["local"] for r in rows})
    statuses = sorted({r["status"] for r in rows})
    profiles = sorted({r["profile"] for r in rows})

    users_by_login, users_by_fullname = build_existing_user_maps(users)

    existing_user_count = 0
    new_user_count = 0
    login_change_count = 0
    already_senior_login_count = 0
    new_logins: List[Tuple[str, str]] = []
    login_changes: List[Tuple[str, str, str]] = []

    for row in rows:
        user = resolve_user_for_row(
            row,
            users_by_login,
            users_by_fullname,
        )

        if user:
            existing_user_count += 1
            current_login = clean_text(user.get("name"))
            desired_login = row["login"]

            if normalize_compare(current_login) != normalize_compare(desired_login):
                login_change_count += 1
                login_changes.append(
                    (row["full_name"], current_login, desired_login)
                )
            else:
                already_senior_login_count += 1
        else:
            new_user_count += 1
            new_logins.append((row["full_name"], row["login"]))

    print()
    print("=" * 96)
    print("ANÁLISE DA PLANILHA")
    print("=" * 96)
    print(f"Linhas válidas para importação........: {len(rows)}")
    print(f"Linhas ignoradas......................: {len(skipped)}")
    print(f"Grupos / Gerências únicas............: {len(gerencias)}")
    print(f"Locais únicos.........................: {len(locais)}")
    print(f"Status únicos.........................: {len(statuses)}")
    print(f"Usuários já existentes pelo nome.....: {existing_user_count}")
    print(f"Usuários que seriam criados..........: {new_user_count}")
    print(f"Logins que serão migrados p/ Senior..: {login_change_count}")
    print(f"Logins já usando Matrícula Senior....: {already_senior_login_count}")
    print(f"Usuários que terão senha atualizada..: {existing_user_count}")
    print(
        "Senha padrão..........................: "
        + ("CONFIGURADA" if DEFAULT_PASSWORD else "NÃO CONFIGURADA")
    )
    print()

    print("Perfis usados na planilha:")
    for profile in profiles:
        pid = profiles_by_name.get(normalize_compare(profile))
        print(f"  - {profile}: ID {pid}")

    print()
    print("Grupos / Gerencia Lotação:")
    for name in gerencias:
        exists = normalize_compare(name) in groups_map
        dd_exists = (
            normalize_compare(name)
            in dropdown_maps["gerencia"]
        )
        print(
            f"  - {name} | grupo={'EXISTE' if exists else 'CRIAR'} "
            f"| dropdown={'EXISTE' if dd_exists else 'CRIAR'}"
        )

    print()
    print("Localização fisica Gerencia:")
    for name in locais:
        exists = normalize_compare(name) in dropdown_maps["local"]
        print(f"  - {name}: {'EXISTE' if exists else 'CRIAR'}")

    print()
    print("Status Mobilização:")
    for name in statuses:
        exists = normalize_compare(name) in dropdown_maps["status"]
        print(f"  - {name}: {'EXISTE' if exists else 'CRIAR'}")

    print()
    print("Container Fields:")
    print(
        f"  Container: {container.get('label')} "
        f"(ID {container.get('id')}, name={container.get('name')})"
    )
    for key in ("status", "gerencia", "local"):
        f = fields[key]
        print(
            f"  {key}: label='{f.get('label')}' "
            f"name='{f.get('name')}' type='{f.get('type')}'"
        )

    print()
    print(
        f"Itemtype de armazenamento do container: "
        f"{instance_itemtype}"
    )
    print(
        f"Registros atuais acessíveis via API nesse itemtype: "
        f"{len(instance_api_rows)}"
    )

    print()
    print("Amostra das migrações de login:")
    for full_name, old_login, new_login in login_changes[:15]:
        print(f"  {full_name}: {old_login} -> {new_login}")

    if len(login_changes) > 15:
        print(
            f"  ... mais {len(login_changes) - 15} usuário(s)"
        )

    if new_logins:
        print()
        print("Novos usuários que seriam criados:")
        for full_name, login in new_logins[:15]:
            print(f"  {full_name} -> {login}")

    print()
    print("Comentários dos usuários:")
    print(
        "  ITEM PPU: <coluna G> :: PREPOSTO: <coluna E>"
    )

    print()
    print("Campos NÃO preenchidos nesta etapa:")
    print("  - Preposto (campo dinâmico)")
    print("  - Chave Colaborador")
    print("  - Líder")

    print()
    print("Linhas ignoradas por 'não mobilizar de imediato':")
    print("  " + ", ".join(str(x) for x in skipped) if skipped else "  nenhuma")

    print("=" * 96)
    print("DRY-RUN concluído. Nenhuma alteração foi realizada.")
    print("=" * 96)


def run_apply(
    glpi: GLPI,
    rows: List[Dict[str, str]],
    profiles_by_name: Dict[str, int],
    groups_map: Dict[str, Dict[str, Any]],
    users: List[Dict[str, Any]],
    profile_users: List[Dict[str, Any]],
    group_users: List[Dict[str, Any]],
    container: Dict[str, Any],
    fields: Dict[str, Dict[str, Any]],
    dropdown_itemtypes: Dict[str, str],
    dropdown_maps: Dict[str, Dict[str, Dict[str, Any]]],
    instance_itemtype: str,
    instance_rows: List[Dict[str, Any]],
) -> None:
    users_by_login, users_by_fullname = build_existing_user_maps(users)

    profile_user_existing = {
        (
            int(x.get("users_id") or 0),
            int(x.get("profiles_id") or 0),
            int(x.get("entities_id") or 0),
            int(x.get("is_recursive") or 0),
        )
        for x in profile_users
    }

    group_user_existing = {
        (
            int(x.get("users_id") or 0),
            int(x.get("groups_id") or 0),
        )
        for x in group_users
    }

    instance_by_user_id = {
        int(x.get("items_id") or 0): x
        for x in instance_rows
        if int(x.get("items_id") or 0) > 0
    }

    gerencias = sorted({r["gerencia"] for r in rows})
    locais = sorted({r["local"] for r in rows})
    statuses = sorted({r["status"] for r in rows})

    group_ids: Dict[str, int] = {}
    for name in gerencias:
        group_ids[name] = ensure_group(
            glpi,
            groups_map,
            name,
        )

    dropdown_ids: Dict[str, Dict[str, int]] = {
        "gerencia": {},
        "local": {},
        "status": {},
    }

    for name in gerencias:
        dropdown_ids["gerencia"][name] = ensure_dropdown_value(
            glpi,
            dropdown_itemtypes["gerencia"],
            dropdown_maps["gerencia"],
            name,
        )

    for name in locais:
        dropdown_ids["local"][name] = ensure_dropdown_value(
            glpi,
            dropdown_itemtypes["local"],
            dropdown_maps["local"],
            name,
        )

    for name in statuses:
        dropdown_ids["status"][name] = ensure_dropdown_value(
            glpi,
            dropdown_itemtypes["status"],
            dropdown_maps["status"],
            name,
        )

    field_names = {
        key: clean_text(field["name"])
        for key, field in fields.items()
    }

    col_status = (
        f"plugin_fields_{field_names['status']}dropdowns_id"
    )
    col_gerencia = (
        f"plugin_fields_{field_names['gerencia']}dropdowns_id"
    )
    col_local = (
        f"plugin_fields_{field_names['local']}dropdowns_id"
    )

    container_id = int(container["id"])

    created_users = 0
    reused_users = 0
    migrated_logins = 0
    updated_fields = 0

    for pos, row in enumerate(rows, start=1):
        full_name = row["full_name"]
        firstname, realname = split_name(full_name)
        fullname_key = normalize_compare(full_name)
        desired_login = row["login"]

        user = resolve_user_for_row(
            row,
            users_by_login,
            users_by_fullname,
        )
        existed_before = user is not None

        if existed_before:
            user_id = int(user["id"])
            old_login = clean_text(user.get("name"))
            login = desired_login
            reused_users += 1
        else:
            login = desired_login

            comment = (
                f"ITEM PPU: {row['ppu']} :: "
                f"PREPOSTO: {row['preposto']}"
            )

            user_id = glpi.add(
                "User",
                {
                    "name": login,
                    "firstname": firstname,
                    "realname": realname,
                    "is_active": 1,
                    "comment": comment,
                    "password": DEFAULT_PASSWORD,
                    "password2": DEFAULT_PASSWORD,
                },
            )

            user = {
                "id": user_id,
                "name": login,
                "firstname": firstname,
                "realname": realname,
            }

            users_by_login[normalize_compare(login)] = user
            users_by_fullname.setdefault(
                fullname_key,
                [],
            ).append(user)

            created_users += 1
            print(
                f"CRIADO usuário ID={user_id}: "
                f"{full_name} -> {login}"
            )

        profile_id = profiles_by_name[
            normalize_compare(row["profile"])
        ]
        group_id = group_ids[row["gerencia"]]

        ensure_profile_user(
            glpi,
            profile_user_existing,
            user_id,
            profile_id,
        )
        ensure_group_user(
            glpi,
            group_user_existing,
            user_id,
            group_id,
        )

        comment = (
            f"ITEM PPU: {row['ppu']} :: "
            f"PREPOSTO: {row['preposto']}"
        )

        # Depois das associações, define perfil, entidade e grupo padrão.
        # Para usuários que já existiam antes desta execução, redefine também
        # a senha para a senha inicial padrão configurada em GLPI_DEFAULT_PASSWORD.
        user_update = {
            "name": login,
            "firstname": firstname,
            "realname": realname,
            "is_active": 1,
            "comment": comment,
            "profiles_id": profile_id,
            "entities_id": ROOT_ENTITY_ID,
            "groups_id": group_id,
        }

        if existed_before:
            user_update["password"] = DEFAULT_PASSWORD
            user_update["password2"] = DEFAULT_PASSWORD

        glpi.update(
            "User",
            user_id,
            user_update,
        )

        if existed_before:
            if normalize_compare(old_login) != normalize_compare(login):
                users_by_login.pop(normalize_compare(old_login), None)
                user["name"] = login
                users_by_login[normalize_compare(login)] = user
                migrated_logins += 1
                print(
                    f"ALTERADO login ID={user_id}: "
                    f"{old_login} -> {login}"
                )

        plugin_payload = {
            "items_id": user_id,
            "itemtype": "User",
            "plugin_fields_containers_id": container_id,
            col_status: dropdown_ids["status"][row["status"]],
            col_gerencia: dropdown_ids["gerencia"][row["gerencia"]],
            col_local: dropdown_ids["local"][row["local"]],
        }

        existing_instance = instance_by_user_id.get(user_id)

        if existing_instance:
            instance_id = int(existing_instance["id"])
            glpi.update(
                instance_itemtype,
                instance_id,
                plugin_payload,
            )
        else:
            instance_id = glpi.add(
                instance_itemtype,
                plugin_payload,
            )
            instance_by_user_id[user_id] = {
                "id": instance_id,
                **plugin_payload,
            }

        updated_fields += 1

        print(
            f"[{pos:03}/{len(rows)}] OK "
            f"ID={user_id} login={login} "
            f"grupo='{row['gerencia']}' "
            f"local='{row['local']}' "
            f"status='{row['status']}' "
            f"perfil='{row['profile']}'"
        )

    print()
    print("=" * 96)
    print("IMPORTAÇÃO CONCLUÍDA")
    print("=" * 96)
    print(f"Usuários criados.....................: {created_users}")
    print(f"Usuários reaproveitados..............: {reused_users}")
    print(f"Logins migrados para Matrícula Senior: {migrated_logins}")
    print(f"Usuários com campos dinâmicos gravados: {updated_fields}")
    print(f"Grupos processados...................: {len(group_ids)}")
    print(f"Gerências dropdown...................: {len(dropdown_ids['gerencia'])}")
    print(f"Locais dropdown......................: {len(dropdown_ids['local'])}")
    print(f"Status dropdown......................: {len(dropdown_ids['status'])}")
    print("=" * 96)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Importa Base 4711 no GLPI."
    )

    parser.add_argument(
        "--xlsx",
        default=DEFAULT_XLSX,
        help=f"Caminho do XLSX. Padrão: {DEFAULT_XLSX}",
    )

    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="Somente valida e apresenta o plano.",
    )
    mode.add_argument(
        "--apply",
        action="store_true",
        help="Executa a importação.",
    )

    args = parser.parse_args()

    if args.apply and not DEFAULT_PASSWORD:
        raise RuntimeError(
            "GLPI_DEFAULT_PASSWORD não definido. "
            "Defina a senha padrão em /root/.glpi.env antes do --apply."
        )

    rows, skipped = read_rows(args.xlsx)

    print(f"OK - planilha lida: {args.xlsx}")
    print(f"OK - {len(rows)} linha(s) válida(s)")
    print(f"OK - {len(skipped)} linha(s) ignorada(s)")

    if len(rows) != 120:
        print(
            f"ATENÇÃO: esperávamos 120 linhas válidas, "
            f"mas foram encontradas {len(rows)}."
        )

    glpi = GLPI()

    try:
        glpi.init()

        profiles = glpi.get_all("Profile")
        profiles_by_name = get_profiles_by_name(profiles)

        required_profiles = sorted({r["profile"] for r in rows})

        missing_profiles = [
            p
            for p in required_profiles
            if normalize_compare(p) not in profiles_by_name
        ]

        if missing_profiles:
            raise RuntimeError(
                "Perfis da planilha não encontrados no GLPI: "
                + ", ".join(missing_profiles)
            )

        print(
            "OK - perfis localizados: "
            + ", ".join(
                f"{p}=ID {profiles_by_name[normalize_compare(p)]}"
                for p in required_profiles
            )
        )

        container, fields = find_container_and_fields(glpi)

        print(
            f"OK - container '{container.get('label')}' "
            f"ID={container.get('id')}"
        )

        dropdown_itemtypes = {
            key: dropdown_itemtype(clean_text(field["name"]))
            for key, field in fields.items()
        }

        dropdown_maps: Dict[
            str,
            Dict[str, Dict[str, Any]],
        ] = {}

        for key, itemtype in dropdown_itemtypes.items():
            rows_dd = glpi.get_all(itemtype)
            dropdown_maps[key] = map_dropdown(rows_dd)
            print(
                f"OK - {itemtype}: "
                f"{len(rows_dd)} valor(es) existente(s)"
            )

        container_name = clean_text(container.get("name"))
        instance_itemtype = container_instance_itemtype(
            container_name
        )

        # Este teste é importante no GLPI 10:
        # verifica se o itemtype gerado pelo Fields está exposto
        # na Legacy REST API.
        instance_rows = glpi.get_all(
            instance_itemtype,
            allow_missing=True,
        )

        # Se retornou vazio pode ser tabela realmente vazia.
        # Verificamos o endpoint diretamente para diferenciar 404/400.
        endpoint_test = glpi.http.get(
            f"{API_URL}/{instance_itemtype}/",
            params={"range": "0-0", "get_hateoas": "false"},
            timeout=30,
        )

        if endpoint_test.status_code not in (200, 206):
            raise RuntimeError(
                f"O itemtype gerado '{instance_itemtype}' não está "
                f"acessível pela API (HTTP {endpoint_test.status_code}). "
                "Neste ambiente GLPI 10 será necessário usar o bridge "
                "PHP/local para gravar os campos do plugin Fields."
            )

        groups = glpi.get_all("Group")
        groups_map = map_existing_groups(groups)

        users = glpi.get_all("User")

        if args.dry_run:
            dry_run_report(
                rows,
                skipped,
                profiles_by_name,
                groups_map,
                users,
                dropdown_maps,
                container,
                fields,
                instance_itemtype,
                instance_rows,
            )
            return 0

        profile_users = glpi.get_all("Profile_User")
        group_users = glpi.get_all("Group_User")

        run_apply(
            glpi,
            rows,
            profiles_by_name,
            groups_map,
            users,
            profile_users,
            group_users,
            container,
            fields,
            dropdown_itemtypes,
            dropdown_maps,
            instance_itemtype,
            instance_rows,
        )

        return 0

    finally:
        glpi.close()


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nExecução interrompida.", file=sys.stderr)
        sys.exit(130)
    except requests.RequestException as exc:
        print(f"\nERRO HTTP: {exc}", file=sys.stderr)
        sys.exit(2)
    except Exception as exc:
        print(f"\nERRO: {exc}", file=sys.stderr)
        sys.exit(1)